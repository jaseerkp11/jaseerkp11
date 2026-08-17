#!/usr/bin/env python3
"""Generate a scannable, recruiter-friendly CV for Mohammed Jaseer (DOCX + PDF)."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor, Twips, Emu
from reportlab.lib.colors import HexColor, white
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Flowable,
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

OUT_DIR = Path("/workspace")

# Palette — professional navy + gold, easy to scan
NAVY = RGBColor(0x0F, 0x2C, 0x3C)
GOLD_RGB = RGBColor(0xC4, 0xA3, 0x5A)
NAVY_HEX = HexColor("#0F2C3C")
NAVY_MID = HexColor("#17384A")
GOLD = HexColor("#C4A35A")
SAND = HexColor("#F6F1E6")
PALE = HexColor("#EEF4F7")
GRAY = HexColor("#2C2C2C")
MUTED = HexColor("#5A6570")
RULE = HexColor("#D9D2C5")


SUMMARY = (
    "Versatile Accounting and Operations professional with UAE experience across accounting, "
    "finance, procurement, purchasing, project coordination, client management, staff supervision, "
    "and business operations in the interior fit-out industry. Joined Introyale Interiors LLC "
    "(Abu Dhabi) as an Accountant and, within the first 5–6 months, took on additional "
    "responsibilities across company operations, procurement, project coordination, client "
    "management, staff supervision, and HR-related activities. Handled the complete workflow from "
    "client RFQ and quotation preparation through LPO receipt, material sourcing, purchasing, "
    "project scheduling, workforce coordination, and client follow-up. Worked directly with major "
    "hotel clients in the UAE and coordinated with purchasing, housekeeping, finance, operations, "
    "and engineering teams. Managed daily work allocation and scheduling for approximately 25–30 "
    "staff. B.Com graduate seeking Accounting, Finance, Procurement, Operations, Project "
    "Coordination, Administration, and related roles in the UAE."
)

SKILLS = [
    (
        "Accounting & Finance",
        "Accounting and Bookkeeping; Accounts Payable (AP); Accounts Receivable (AR); "
        "Purchase and Sales Entries; Invoicing and Billing; Receivables Follow-up; "
        "Payment Follow-up; Financial Documentation; VAT Preparation and Filing",
    ),
    (
        "Procurement & Purchasing",
        "Procurement and Purchasing; Supplier Sourcing; Supplier Coordination; "
        "Supplier Negotiation; Quotation Comparison; Material Sourcing; Purchase Coordination; "
        "Cost Calculation; Cost Control; Delivery Coordination",
    ),
    (
        "Project & Operations",
        "Project Initiation; Project Coordination; Project Scheduling; Daily Work Planning; "
        "Workforce Coordination; Staff Supervision; Task Allocation; Deadline Management; "
        "Material Planning; Operational Coordination",
    ),
    (
        "Client & Business Coordination",
        "RFQ Management; Quotation Preparation; Pricing; LPO Processing; Client Requirement "
        "Gathering; Client Communication; Hotel Client Coordination; Meeting and Site Visit "
        "Coordination; Cross-functional Coordination; Project Documentation",
    ),
]

ACCOUNTING_BULLETS = [
    "Managed day-to-day accounting activities including purchase entries, sales entries, accounts payable (AP), and accounts receivable (AR).",
    "Recorded and maintained financial transactions and supporting documentation required for financial and operational activities.",
    "Prepared invoices, maintained billing records, monitored customer outstanding balances, and followed up with clients on receivables by email and telephone.",
    "Prepared VAT-related records and handled VAT preparation and filing activities.",
    "Coordinated financial follow-ups connected with ongoing business and project activities.",
]

RFQ_BULLETS = [
    "Managed the quotation workflow from client RFQ receipt through requirement review, site visits, hotel client meetings, quotation submission, follow-up, and approval.",
    "Prepared quotations and calculated project pricing based on requirements, materials, labour, and related costs; coordinated with purchasing teams and suppliers for material information and pricing.",
    "Submitted quotations, followed up on clarifications and approvals, maintained client communication, processed approved LPOs, and initiated procurement and project execution.",
]

PROCUREMENT_BULLETS = [
    "Managed material sourcing after receiving approved client LPOs and identified suppliers according to project requirements.",
    "Obtained and compared supplier quotations; negotiated purchase prices and coordinated procurement according to project budgets and deadlines.",
    "Calculated project-related costs, monitored expenses, and worked to minimize unnecessary expenditure while meeting project requirements.",
    "Coordinated purchasing and material delivery schedules with project deadlines and followed up with suppliers on availability, purchasing, and delivery.",
]

PROJECT_BULLETS = [
    "Coordinated interior fit-out projects from initiation through execution, prioritizing approved LPOs according to client deadlines.",
    "Prepared project and workforce schedules; coordinated materials, staff, suppliers, and internal teams; and monitored ongoing and upcoming project requirements.",
    "Communicated with clients on progress, requirements, approvals, and schedules; attended meetings and coordinated between hotel representatives and internal teams.",
    "Coordinated with purchasing managers, housekeeping managers, finance, operations, and engineering departments.",
    "Prepared daily work schedules, assigned tasks and deadlines, followed up on progress, and coordinated staff according to project priorities.",
]

STAFF_BULLETS = [
    "Expanded from Accountant into company-wide operations within approximately 5–6 months, handling day-to-day operational coordination alongside accounting.",
    "Supervised and coordinated approximately 25–30 staff, including daily schedules, work assignment, deadline communication, and workforce planning.",
    "Supported HR-related staff management activities and acted as a key coordination point between management, staff, clients, suppliers, and project stakeholders.",
]

TUBEES_BULLETS = [
    "Assisted with recording daily accounting transactions and maintaining financial records and accounting documentation.",
    "Supported invoice preparation, payment follow-ups, and routine accounting and administrative activities.",
    "Supported general office and operational requirements.",
]


NAVY_FILL = "0F2C3C"
GOLD_FILL = "C4A35A"
SAND_FILL = "F6F1E6"
PALE_FILL = "EEF4F7"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x2C, 0x2C, 0x2C)
MUTED_RGB = RGBColor(0x5A, 0x65, 0x70)
DATE_GOLD = RGBColor(0x8A, 0x70, 0x40)


def set_run_font(run, name="Calibri", size=11, bold=False, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_spacing(p, before=0, after=0, line=230):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = Twips(line)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for m, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def set_cell_borders(cell, **edges):
    """edges like top={'sz':'18','color':'C4A35A','val':'single'}."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        spec = edges.get(edge, {"sz": "0", "val": "nil", "color": "auto"})
        el.set(qn("w:val"), spec.get("val", "single"))
        el.set(qn("w:sz"), spec.get("sz", "0"))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), spec.get("color", "auto"))
        tc_borders.append(el)
    tc_pr.append(tc_borders)


def set_table_full_width(table):
    table.autofit = True
    table.allow_autofit = True
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)


def clear_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tbl_pr.append(borders)


def cell_text(cell, text, size=10, bold=False, color=TEXT, align="left", before=0, after=0, italic=False):
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    set_paragraph_spacing(p, before=before, after=after, line=220)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_cell_line(cell, text, size=10, bold=False, color=TEXT, align="left", before=0, after=0, italic=False):
    p = cell.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    set_paragraph_spacing(p, before=before, after=after, line=220)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def tight_para(doc, before=0, after=0):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after, line=200)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    return p


def add_section_bar(doc, title):
    tight_para(doc, before=6, after=2)
    table = doc.add_table(rows=1, cols=1)
    set_table_full_width(table)
    clear_table_borders(table)
    cell = table.cell(0, 0)
    shade_cell(cell, NAVY_FILL)
    set_cell_margins(cell, top=50, bottom=50, left=120, right=80)
    set_cell_borders(
        cell,
        left={"sz": "24", "val": "single", "color": GOLD_FILL},
        top={"sz": "0", "val": "nil", "color": "auto"},
        bottom={"sz": "0", "val": "nil", "color": "auto"},
        right={"sz": "0", "val": "nil", "color": "auto"},
    )
    cell_text(cell, title.upper(), size=10, bold=True, color=WHITE, align="left")
    tight_para(doc, before=2, after=2)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=1, line=216)
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    run = p.add_run("•  " + text)
    set_run_font(run, size=10, color=TEXT)
    return p


def add_subhead(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=6, after=2, line=220)
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=True, color=NAVY)
    return p


def add_job_header(doc, title, dates, meta=None, industry=None):
    table = doc.add_table(rows=1, cols=2)
    set_table_full_width(table)
    clear_table_borders(table)
    left, right = table.cell(0, 0), table.cell(0, 1)
    set_cell_margins(left, top=40, bottom=20, left=0, right=40)
    set_cell_margins(right, top=40, bottom=20, left=40, right=0)
    gold_line = {"sz": "6", "val": "single", "color": "D9D2C5"}
    for cell in (left, right):
        set_cell_borders(
            cell,
            bottom=gold_line,
            top={"sz": "0", "val": "nil", "color": "auto"},
            left={"sz": "0", "val": "nil", "color": "auto"},
            right={"sz": "0", "val": "nil", "color": "auto"},
        )
    cell_text(left, title, size=11, bold=True, color=NAVY, align="left")
    cell_text(right, dates, size=10, bold=True, color=DATE_GOLD, align="right")
    if meta:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=2, after=0, line=220)
        run = p.add_run(meta)
        set_run_font(run, size=10, bold=True, color=RGBColor(0x3D, 0x5A, 0x68))
    if industry:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=4, line=220)
        run = p.add_run(industry)
        set_run_font(run, size=9.5, italic=True, color=MUTED_RGB)
    return table


def add_page_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, before=4, after=0, line=200)
    # Navy bar look via paragraph shading
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), NAVY_FILL)
    p._p.get_or_add_pPr().append(shd)
    run = p.add_run("  Mohammed Jaseer  |  Curriculum Vitae")
    set_run_font(run, size=8, color=WHITE)
    run = p.add_run("\t")
    set_run_font(run, size=8, color=WHITE)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.2), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run("Page ")
    set_run_font(run, size=8, color=WHITE)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run2 = p.add_run()
    run2._r.append(fld)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run3 = p.add_run()
    run3._r.append(instr)
    set_run_font(run3, size=8, color=WHITE)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run4 = p.add_run()
    run4._r.append(fld_end)


def build_docx(path: Path):
    """Match the PDF layout: navy header, gold rule, snapshot cards, skill tiles."""
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(10)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(14)
    section.right_margin = Mm(14)
    add_page_footer(section)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.font.color.rgb = TEXT

    # Header band
    header = doc.add_table(rows=4, cols=1)
    set_table_full_width(header)
    clear_table_borders(header)
    rows = [
        ("MOHAMMED JASEER", 20, True, WHITE, 8, 0),
        (
            "Accounting  |  Finance  |  Operations  |  Procurement  |  Project Coordination",
            10.5,
            True,
            GOLD_RGB,
            0,
            2,
        ),
        (
            "Currently in India  |  Available for UAE Relocation  |  UAE Residence Visa valid until December 2027",
            9.5,
            False,
            WHITE,
            2,
            0,
        ),
        (
            "+971 56 454 9259  |  mohammedjaseerkp11@gmail.com  |  linkedin.com/in/jaseerkp/",
            10,
            False,
            WHITE,
            0,
            8,
        ),
    ]
    for i, (text, size, bold, color, before, after) in enumerate(rows):
        cell = header.cell(i, 0)
        shade_cell(cell, NAVY_FILL)
        set_cell_margins(cell, top=40 if i else 80, bottom=40 if i < 3 else 90, left=100, right=100)
        set_cell_borders(cell)
        cell_text(cell, text, size=size, bold=bold, color=color, align="center", before=before, after=after)

    # Gold rule
    rule = doc.add_table(rows=1, cols=1)
    set_table_full_width(rule)
    clear_table_borders(rule)
    rc = rule.cell(0, 0)
    shade_cell(rc, GOLD_FILL)
    set_cell_margins(rc, top=0, bottom=0, left=0, right=0)
    set_cell_borders(rc)
    p = rc.paragraphs[0]
    set_paragraph_spacing(p, before=0, after=0, line=80)
    run = p.add_run(" ")
    set_run_font(run, size=4, color=GOLD_RGB)

    tight_para(doc, before=6, after=4)

    # Snapshot cards — same as PDF
    snaps = [
        ("UAE EXPERIENCE", "Abu Dhabi"),
        ("TEAM SUPERVISED", "25–30 Staff"),
        ("RESIDENCE VISA", "Valid to Dec 2027"),
        ("INDUSTRY", "Hotel Fit-Out"),
    ]
    snap = doc.add_table(rows=1, cols=4)
    set_table_full_width(snap)
    clear_table_borders(snap)
    gold_top = {"sz": "18", "val": "single", "color": GOLD_FILL}
    for i, (label, value) in enumerate(snaps):
        cell = snap.cell(0, i)
        shade_cell(cell, SAND_FILL)
        set_cell_margins(cell, top=70, bottom=70, left=50, right=50)
        set_cell_borders(cell, top=gold_top)
        cell_text(cell, label, size=7.5, bold=False, color=MUTED_RGB, align="center", after=0)
        add_cell_line(cell, value, size=10, bold=True, color=NAVY, align="center", before=0, after=2)

    add_section_bar(doc, "Professional Summary")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, before=0, after=4, line=220)
    run = p.add_run(SUMMARY)
    set_run_font(run, size=10, color=TEXT)

    add_section_bar(doc, "Core Skills")
    skills_tbl = doc.add_table(rows=2, cols=2)
    set_table_full_width(skills_tbl)
    clear_table_borders(skills_tbl)
    gold_left = {"sz": "24", "val": "single", "color": GOLD_FILL}
    for i, (title, body) in enumerate(SKILLS):
        cell = skills_tbl.cell(i // 2, i % 2)
        shade_cell(cell, PALE_FILL)
        set_cell_margins(cell, top=70, bottom=70, left=90, right=70)
        set_cell_borders(cell, left=gold_left)
        cell_text(cell, title, size=10, bold=True, color=NAVY, align="left", after=2)
        add_cell_line(cell, body, size=9, bold=False, color=TEXT, align="left", after=2)

    add_section_bar(doc, "Professional Experience")
    add_job_header(
        doc,
        "INTROYALE INTERIORS LLC — Abu Dhabi, UAE",
        "27 Oct 2023 – 16 May 2026",
        meta="Accountant → Expanded Operations, Procurement & Project Responsibilities",
        industry="Industry: Interior Fit-Out / Hotel Projects",
    )
    add_subhead(doc, "Accounting & Finance")
    for t in ACCOUNTING_BULLETS:
        add_bullet(doc, t)
    add_subhead(doc, "RFQ, Quotation & Client Coordination")
    for t in RFQ_BULLETS:
        add_bullet(doc, t)
    add_subhead(doc, "Procurement & Purchasing")
    for t in PROCUREMENT_BULLETS:
        add_bullet(doc, t)
    add_subhead(doc, "Project Coordination & Execution")
    for t in PROJECT_BULLETS:
        add_bullet(doc, t)
    add_subhead(doc, "Staff & Operations Management")
    for t in STAFF_BULLETS:
        add_bullet(doc, t)

    tight_para(doc, before=4, after=2)
    add_job_header(doc, "TUBEES FOODS — Junior Accounts Assistant", "2022 – 2023")
    for t in TUBEES_BULLETS:
        add_bullet(doc, t)

    add_section_bar(doc, "Education")
    add_job_header(doc, "UNIVERSITY OF CALICUT — Bachelor of Commerce (B.Com)", "2019 – 2021")
    add_job_header(doc, "HSS ANAMANGAD — Higher Secondary Education", "2017 – 2018")

    add_section_bar(doc, "Languages")
    langs = doc.add_table(rows=1, cols=3)
    set_table_full_width(langs)
    clear_table_borders(langs)
    for i, (name, level) in enumerate(
        [("Malayalam", "Native"), ("English", "Fluent"), ("Hindi", "Basic")]
    ):
        cell = langs.cell(0, i)
        shade_cell(cell, SAND_FILL)
        set_cell_margins(cell, top=70, bottom=70, left=40, right=40)
        set_cell_borders(cell, top=gold_top)
        cell_text(cell, name, size=11, bold=True, color=NAVY, align="center", after=0)
        add_cell_line(cell, level, size=9, bold=False, color=MUTED_RGB, align="center", after=2)

    add_section_bar(doc, "Additional Information")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=2, after=2, line=220)
    run = p.add_run(
        "Current Location: India  |  UAE Experience: Abu Dhabi  |  "
        "UAE Residence Visa: Valid until December 2027  |  "
        "Relocation: Available for UAE opportunities  |  "
        "Availability: Open to suitable opportunities"
    )
    set_run_font(run, size=9.5, color=TEXT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=2, line=220)
    run = p.add_run("Target Roles: ")
    set_run_font(run, size=9.5, bold=True, color=TEXT)
    run = p.add_run(
        "Accounting | Finance | Procurement | Purchasing | "
        "Operations | Project Coordination | Administration | Business Support"
    )
    set_run_font(run, size=9.5, color=TEXT)

    doc.save(path)


def html_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def para(text, style):
    return Paragraph(text, style)


class GoldRule(Flowable):
    def __init__(self, width, thickness=1.6):
        super().__init__()
        self.width = width
        self.height = thickness
        self.thickness = thickness

    def draw(self):
        self.canv.setFillColor(GOLD)
        self.canv.rect(0, 0, self.width, self.thickness, fill=1, stroke=0)


class SectionBar(Flowable):
    def __init__(self, title, width):
        super().__init__()
        self.title = title.upper()
        self.width = width
        self.height = 16

    def wrap(self, availWidth, availHeight):
        self.width = availWidth
        return self.width, self.height

    def draw(self):
        c = self.canv
        c.setFillColor(NAVY_HEX)
        c.roundRect(0, 0, self.width, 15, 1.5, fill=1, stroke=0)
        c.setFillColor(GOLD)
        c.rect(0, 0, 4, 15, fill=1, stroke=0)
        c.setFillColor(white)
        c.setFont("Helvetica-Bold", 9)
        c.drawString(10, 4.5, self.title)


def build_pdf(path: Path):
    page_w, page_h = A4
    side = 14 * mm
    content_w = page_w - 2 * side

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=side,
        rightMargin=side,
        topMargin=11 * mm,
        bottomMargin=12 * mm,
        title="Mohammed Jaseer — CV",
        author="Mohammed Jaseer",
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="Name",
            fontName="Helvetica-Bold",
            fontSize=20,
            leading=24,
            alignment=TA_CENTER,
            textColor=white,
            spaceAfter=1,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Headline",
            fontName="Helvetica-Bold",
            fontSize=9.2,
            leading=12,
            alignment=TA_CENTER,
            textColor=GOLD,
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Contact",
            fontName="Helvetica",
            fontSize=8.6,
            leading=11.2,
            alignment=TA_CENTER,
            textColor=white,
            spaceAfter=0,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SnapLabel",
            fontName="Helvetica",
            fontSize=7.2,
            leading=9,
            alignment=TA_CENTER,
            textColor=MUTED,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SnapValue",
            fontName="Helvetica-Bold",
            fontSize=8.4,
            leading=11,
            alignment=TA_CENTER,
            textColor=NAVY_HEX,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BodyJust",
            fontName="Helvetica",
            fontSize=9.1,
            leading=12,
            alignment=TA_JUSTIFY,
            textColor=GRAY,
            spaceAfter=3,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SkillTitle",
            fontName="Helvetica-Bold",
            fontSize=8.4,
            leading=11,
            textColor=NAVY_HEX,
            spaceAfter=2,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SkillBody",
            fontName="Helvetica",
            fontSize=8.1,
            leading=10.6,
            textColor=GRAY,
        )
    )
    styles.add(
        ParagraphStyle(
            name="JobTitle",
            fontName="Helvetica-Bold",
            fontSize=9.6,
            leading=12.4,
            textColor=NAVY_HEX,
        )
    )
    styles.add(
        ParagraphStyle(
            name="JobMeta",
            fontName="Helvetica-Bold",
            fontSize=8.8,
            leading=11.5,
            textColor=HexColor("#3D5A68"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="JobDates",
            fontName="Helvetica-Bold",
            fontSize=8.6,
            leading=12.4,
            alignment=TA_RIGHT,
            textColor=HexColor("#8A7040"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="SubHead",
            fontName="Helvetica-Bold",
            fontSize=8.8,
            leading=11.5,
            textColor=NAVY_HEX,
            spaceBefore=3.5,
            spaceAfter=1.2,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BulletText",
            fontName="Helvetica",
            fontSize=8.85,
            leading=11.4,
            textColor=GRAY,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Industry",
            fontName="Helvetica-Oblique",
            fontSize=8.3,
            leading=10.8,
            textColor=MUTED,
            spaceAfter=1,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Lang",
            fontName="Helvetica",
            fontSize=8.8,
            leading=11.5,
            alignment=TA_CENTER,
            textColor=GRAY,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SmallCenter",
            fontName="Helvetica",
            fontSize=8.3,
            leading=11,
            alignment=TA_CENTER,
            textColor=GRAY,
        )
    )

    story = []

    def section(title):
        story.append(Spacer(1, 5))
        story.append(SectionBar(title, content_w))
        story.append(Spacer(1, 5))

    def bullets(items):
        for item in items:
            story.append(
                para("&#8226;&nbsp;&nbsp;" + html_escape(item), styles["BulletText"])
            )
            story.append(Spacer(1, 0.6))

    def job_block(left, right, meta=None, industry=None):
        t = Table(
            [[para(left, styles["JobTitle"]), para(right, styles["JobDates"])]],
            colWidths=[content_w * 0.72, content_w * 0.28],
        )
        t.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.4, RULE),
                ]
            )
        )
        story.append(t)
        if meta:
            story.append(para(html_escape(meta), styles["JobMeta"]))
        if industry:
            story.append(para(html_escape(industry), styles["Industry"]))

    # Header band
    header = Table(
        [
            [para("MOHAMMED JASEER", styles["Name"])],
            [
                para(
                    html_escape(
                        "Accounting  |  Finance  |  Operations  |  Procurement  |  Project Coordination"
                    ),
                    styles["Headline"],
                )
            ],
            [
                para(
                    html_escape(
                        "Currently in India  |  Available for UAE Relocation  |  UAE Residence Visa valid until December 2027"
                    ),
                    styles["Contact"],
                )
            ],
            [
                para(
                    html_escape(
                        "+971 56 454 9259  |  mohammedjaseerkp11@gmail.com  |  linkedin.com/in/jaseerkp/"
                    ),
                    styles["Contact"],
                )
            ],
        ],
        colWidths=[content_w],
    )
    header.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), NAVY_HEX),
                ("TOPPADDING", (0, 0), (-1, 0), 10),
                ("BOTTOMPADDING", (0, -1), (-1, -1), 10),
                ("TOPPADDING", (0, 1), (-1, -1), 1),
                ("BOTTOMPADDING", (0, 0), (-1, -2), 1),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(header)
    story.append(GoldRule(content_w, 3))
    story.append(Spacer(1, 6))

    # Snapshot cards — same facts, easier to notice in 5 seconds
    snaps = [
        ("UAE EXPERIENCE", "Abu Dhabi"),
        ("TEAM SUPERVISED", "25–30 Staff"),
        ("RESIDENCE VISA", "Valid to Dec 2027"),
        ("INDUSTRY", "Hotel Fit-Out"),
    ]
    snap_cells = []
    for label, value in snaps:
        inner = Table(
            [
                [para(label, styles["SnapLabel"])],
                [para(html_escape(value), styles["SnapValue"])],
            ],
            colWidths=[content_w / 4 - 6],
        )
        inner.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), SAND),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("TOPPADDING", (0, 0), (-1, 0), 5),
                    ("BOTTOMPADDING", (0, -1), (-1, -1), 6),
                    ("TOPPADDING", (0, 1), (-1, 1), 0),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("LINEABOVE", (0, 0), (-1, 0), 2.2, GOLD),
                ]
            )
        )
        snap_cells.append(inner)
    snap_row = Table([snap_cells], colWidths=[content_w / 4] * 4)
    snap_row.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2.5),
            ]
        )
    )
    story.append(snap_row)

    section("Professional Summary")
    story.append(para(html_escape(SUMMARY), styles["BodyJust"]))

    section("Core Skills")
    skill_cells = []
    for title, body in SKILLS:
        box = Table(
            [
                [para(html_escape(title), styles["SkillTitle"])],
                [para(html_escape(body), styles["SkillBody"])],
            ],
            colWidths=[content_w / 2 - 6],
        )
        box.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), PALE),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                    ("TOPPADDING", (0, 0), (-1, 0), 6),
                    ("BOTTOMPADDING", (0, -1), (-1, -1), 6),
                    ("TOPPADDING", (0, 1), (-1, 1), 0),
                    ("LINEBEFORE", (0, 0), (0, -1), 3, GOLD),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        skill_cells.append(box)
    skill_grid = Table(
        [[skill_cells[0], skill_cells[1]], [skill_cells[2], skill_cells[3]]],
        colWidths=[content_w / 2, content_w / 2],
    )
    skill_grid.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2.5),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ]
        )
    )
    story.append(skill_grid)

    section("Professional Experience")
    job_block(
        html_escape("INTROYALE INTERIORS LLC — Abu Dhabi, UAE"),
        html_escape("27 Oct 2023 – 16 May 2026"),
        meta="Accountant → Expanded Operations, Procurement & Project Responsibilities",
        industry="Industry: Interior Fit-Out / Hotel Projects",
    )

    story.append(para("Accounting &amp; Finance", styles["SubHead"]))
    bullets(ACCOUNTING_BULLETS)
    story.append(para("RFQ, Quotation &amp; Client Coordination", styles["SubHead"]))
    bullets(RFQ_BULLETS)
    story.append(para("Procurement &amp; Purchasing", styles["SubHead"]))
    bullets(PROCUREMENT_BULLETS)
    story.append(para("Project Coordination &amp; Execution", styles["SubHead"]))
    bullets(PROJECT_BULLETS)
    story.append(para("Staff &amp; Operations Management", styles["SubHead"]))
    bullets(STAFF_BULLETS)

    story.append(Spacer(1, 4))
    job_block(
        html_escape("TUBEES FOODS — Junior Accounts Assistant"),
        html_escape("2022 – 2023"),
    )
    story.append(Spacer(1, 2))
    bullets(TUBEES_BULLETS)

    section("Education")
    job_block(
        html_escape("UNIVERSITY OF CALICUT — Bachelor of Commerce (B.Com)"),
        html_escape("2019 – 2021"),
    )
    job_block(
        html_escape("HSS ANAMANGAD — Higher Secondary Education"),
        html_escape("2017 – 2018"),
    )

    section("Languages")
    lang_items = [
        ("Malayalam", "Native"),
        ("English", "Fluent"),
        ("Hindi", "Basic"),
    ]
    lang_cells = []
    for name, level in lang_items:
        box = Table(
            [
                [para(html_escape(name), styles["SnapValue"])],
                [para(html_escape(level), styles["SnapLabel"])],
            ],
            colWidths=[content_w / 3 - 8],
        )
        box.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), SAND),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("TOPPADDING", (0, 0), (-1, 0), 5),
                    ("BOTTOMPADDING", (0, -1), (-1, -1), 5),
                    ("LINEABOVE", (0, 0), (-1, 0), 2, GOLD),
                ]
            )
        )
        lang_cells.append(box)
    langs = Table([lang_cells], colWidths=[content_w / 3] * 3)
    langs.setStyle(
        TableStyle(
            [
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(langs)

    section("Additional Information")
    story.append(
        para(
            html_escape(
                "Current Location: India  |  UAE Experience: Abu Dhabi  |  "
                "UAE Residence Visa: Valid until December 2027  |  "
                "Relocation: Available for UAE opportunities  |  "
                "Availability: Open to suitable opportunities"
            ),
            styles["SmallCenter"],
        )
    )
    story.append(
        para(
            "<b>Target Roles:</b> Accounting | Finance | Procurement | Purchasing | "
            "Operations | Project Coordination | Administration | Business Support",
            styles["SmallCenter"],
        )
    )

    def decorate(canvas, doc_):
        canvas.saveState()
        canvas.setFillColor(NAVY_HEX)
        canvas.rect(0, 0, 4.2 * mm, page_h, fill=1, stroke=0)
        canvas.setFillColor(GOLD)
        canvas.rect(4.2 * mm, 0, 1.1 * mm, page_h, fill=1, stroke=0)
        canvas.setFillColor(NAVY_HEX)
        canvas.rect(0, 0, page_w, 8 * mm, fill=1, stroke=0)
        canvas.setFillColor(GOLD)
        canvas.rect(0, 8 * mm, page_w, 1.1 * mm, fill=1, stroke=0)
        canvas.setFillColor(white)
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(side, 3.2 * mm, "Mohammed Jaseer  |  Curriculum Vitae")
        canvas.drawRightString(page_w - side, 3.2 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)


def main():
    out = OUT_DIR / "cv"
    out.mkdir(exist_ok=True)
    docx_path = out / "Mohammed_Jaseer_CV_ATS.docx"
    pdf_path = out / "Mohammed_Jaseer_CV_ATS.pdf"
    build_docx(docx_path)
    build_pdf(pdf_path)
    print(f"Wrote {docx_path}")
    print(f"Wrote {pdf_path}")


if __name__ == "__main__":
    main()
